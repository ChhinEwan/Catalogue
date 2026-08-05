from datetime import datetime
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
import os
from images_manager import creer_nom_image


def exporter_excel(livres):
    fichier = "bibliotheque.xlsx"
    classeur = Workbook()

    resume = classeur.active
    resume.title = "Résumé"
    resume["A1"] = "Export de la bibliothèque"
    resume["A2"] = f"Date : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    resume["A4"] = "Nombre de livres"
    resume["B4"] = len(livres)

    resume["A6"] = "Genre"
    resume["B6"] = "Nombre"
    resume["A6"].font = Font(bold=True)
    resume["B6"].font = Font(bold=True)

    genres = {}
    for livre in livres:
        genre = livre.genre or "Genre non repertorié"
        genres[genre] = genres.get(genre, 0) + 1

    ligne_resume = 7
    for genre, compte in sorted(genres.items(), key=lambda item: item[1], reverse=True):
        resume[f"A{ligne_resume}"] = genre
        resume[f"B{ligne_resume}"] = compte
        ligne_resume += 1

    feuille = classeur.create_sheet(title="Livres")
    feuille.append([
        "Titre",
        "Auteur",
        "Éditeur",
        "Genre",
        "Date de parution",
        "ISBN",
        "Description",
        "Image"
    ])

    largeurs = [40, 25, 25, 25, 16, 18, 60, 18]
    for index, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur

    for index in range(1, 9):
        cellule = feuille.cell(row=1, column=index)
        cellule.font = Font(bold=True)
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for livre in livres:
        ligne = feuille.max_row + 1
        chemin_image = os.path.join("images", creer_nom_image(livre))

        # prepare truncated description for export (700 chars max)
        description = livre.description or ""
        if len(description) > 700:
            export_description = description[:700].rstrip() + "..."
        else:
            export_description = description or "Description non repertorié"

        feuille.append([
            livre.titre,
            livre.auteur,
            livre.editeur or "Éditeur non repertorié",
            livre.genre or "Genre non repertorié",
            livre.date_parution or "Date inconnu",
            livre.isbn or "ISBN non repertorié",
            export_description,
            ""
        ])

        for colonne in range(1, 9):
            cellule = feuille.cell(row=ligne, column=colonne)
            cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # compute row height from exported (possibly truncated) description
        description_lines = export_description.count("\n") + max(1, (len(export_description) // 50) + 1)
        row_height = description_lines * 15

        if os.path.exists(chemin_image):
            try:
                image = Image(chemin_image)
                image.width = 100
                image.height = 150
                feuille.add_image(image, f"H{ligne}")
                row_height = max(row_height, image.height * 0.75)
            except Exception:
                feuille[f"H{ligne}"] = "img non repertorie"
                feuille[f"H{ligne}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            feuille[f"H{ligne}"] = "img non repertorie"
            feuille[f"H{ligne}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        feuille.row_dimensions[ligne].height = row_height

    try:
        classeur.save(fichier)
    except PermissionError as err:
        raise PermissionError(
            f"Impossible d'enregistrer le fichier '{fichier}'. Fermez-le s'il est ouvert et réessayez."
        ) from err

    return fichier